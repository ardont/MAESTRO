import os
import glob
import sys
import docx
from pptx import Presentation
import fitz # PyMuPDF

target_files = [
    r"C:\Users\Maxim\Downloads\Grazhdanskiy_kodex_Rossiyskoy_Federatsii_chast_pervaya_ot_3 (1).docx",
    r"C:\Users\Maxim\Downloads\Grazhdanskiy_kodex_Rossiyskoy_Federatsii_chast_pervaya_ot_3.docx",
    r"C:\Users\Maxim\Downloads\Federalny_zakon_ot_06_04_2011_N_63_FZ_red_ot_21_04_2025 (1).docx",
    r"C:\Users\Maxim\Downloads\Federalny_zakon_ot_06_04_2011_N_63_FZ_red_ot_21_04_2025.docx",
    r"C:\Users\Maxim\Downloads\Federalny_zakon_ot_18_07_2011_N_223_FZ_red_ot_08_08_2024.docx",
    r"C:\Users\Maxim\Downloads\Federalny_zakon_ot_26_07_2006_N_135_FZ_red_ot_24_06_2025 (1).docx",
    r"C:\Users\Maxim\Downloads\Federalny_zakon_ot_26_07_2006_N_135_FZ_red_ot_24_06_2025.docx",
    r"C:\Users\Maxim\Downloads\03_09_Prezentatsia_dlya_khakatona_6_7_sentyabrya (1).pptx",
    r"C:\Users\Maxim\Downloads\03_09_Prezentatsia_dlya_khakatona_6_7_sentyabrya.pdf",
    r"C:\Users\Maxim\Downloads\Federalny_zakon_ot_26_12_2008_N_294_FZ_red_ot_26_12_2024 (1).docx",
    r"C:\Users\Maxim\Downloads\Federalny_zakon_ot_26_12_2008_N_294_FZ_red_ot_26_12_2024.docx",
    r"C:\Users\Maxim\Downloads\03_09_Prezentatsia_dlya_khakatona_6_7_sentyabrya.pptx",
]

output_dir = r"c:\Users\Maxim\Desktop\фокусы\хакатоны\рлт\агентная система поддержки\txt_dataset\local_files"
os.makedirs(output_dir, exist_ok=True)

def parse_docx(path):
    doc = docx.Document(path)
    lines = []
    # Extract paragraphs and tables in sequence or paragraphs then tables
    for elem in doc.element.body:
        if elem.tag.endswith('p'):
            p = docx.text.paragraph.Paragraph(elem, doc)
            if p.text:
                lines.append(p.text)
        elif elem.tag.endswith('tbl'):
            table = docx.table.Table(elem, doc)
            for row in table.rows:
                row_str = " | ".join(cell.text.strip().replace('\n', ' ') for cell in row.cells)
                if row_str:
                    lines.append(row_str)
    return "\n".join(lines)

def parse_pptx(path):
    prs = Presentation(path)
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                lines.append(shape.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[Notes: {notes}]")
    return "\n".join(lines)

def parse_pdf(path):
    doc = fitz.open(path)
    lines = []
    for i, page in enumerate(doc, 1):
        lines.append(f"--- Page {i} ---")
        lines.append(page.get_text())
    return "\n".join(lines)

results = []
for file_path in target_files:
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        continue
    
    base_name = os.path.basename(file_path)
    name_without_ext = os.path.splitext(base_name)[0]
    out_txt_name = f"{name_without_ext}.txt"
    out_path = os.path.join(output_dir, out_txt_name)
    
    print(f"Processing: {base_name} ...")
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == ".docx":
            content = parse_docx(file_path)
        elif ext == ".pptx":
            content = parse_pptx(file_path)
        elif ext == ".pdf":
            content = parse_pdf(file_path)
        else:
            content = ""
        
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(content)
        
        char_count = len(content)
        line_count = len(content.splitlines())
        print(f"Successfully created {out_txt_name} ({char_count} chars, {line_count} lines)")
        results.append((base_name, out_txt_name, char_count, line_count))
    except Exception as e:
        print(f"Error processing {base_name}: {e}")

print("\n--- SUMMARY ---")
for r in results:
    print(f"{r[0]} -> {r[1]}: {r[2]} chars, {r[3]} lines")
